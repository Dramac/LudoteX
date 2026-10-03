"""
Motif `oubli` : un prêt jamais scanné en retour, clos plus tard par un re-prêt
ou par l'escalade du transfert (lot 13 de la série pré-production, constat
UX-03 ; fiche ouverte au lot agora-1).

Trois choses sont verrouillées ici :

1. le motif est posé par les DEUX chemins — un seul côté rendrait les deux
   gestes incomparables, c'est toute la raison pour laquelle la fiche avait
   été laissée ouverte ;
2. les statistiques le comptent comme un prêt (c'en est un) mais l'écartent
   de la durée moyenne, et la liste détaillée comme les exports disent
   pourquoi sa durée manque ;
3. une base d'avant garde ses lignes telles quelles : aucune migration ne
   requalifie rétroactivement un `pret` déjà clos.
"""

import sqlite3
from io import BytesIO

import pytest

from app import exports, models, services


@pytest.fixture
def conn():
    c = sqlite3.connect(":memory:")
    c.row_factory = sqlite3.Row
    for p in models.PRAGMAS:
        c.execute(p)
    for s in models.SCHEMA_STATEMENTS:
        c.executescript(s)
    c.execute("INSERT INTO titres (reference_titre, nom) VALUES ('CATAN', 'Catan')")
    c.execute("INSERT INTO titres (reference_titre, nom) VALUES ('DIXIT', 'Dixit')")
    c.executemany(
        "INSERT INTO exemplaires (id_exemplaire, reference_titre) VALUES (?, ?)",
        [("001", "CATAN"), ("002", "DIXIT"), ("003", "CATAN")],
    )
    c.commit()
    yield c
    c.close()


def _motifs(conn) -> list[str]:
    return [r[0] for r in conn.execute("SELECT motif FROM prets ORDER BY id_pret")]


# ---------------------------------------------------------------------------
# 1. Les deux chemins
# ---------------------------------------------------------------------------
def test_le_re_pret_pose_le_motif_oubli(conn, vieillir_prets):
    services.preter(conn, "001")
    vieillir_prets(conn)

    services.repreter(conn, "001")

    assert _motifs(conn) == [services.MOTIF_OUBLI, "pret"]


def test_l_escalade_du_transfert_pose_le_motif_oubli(conn, vieillir_prets):
    services.preter(conn, "001")            # le visiteur et sa pochette
    services.preter(conn, "002")            # le prêt oublié
    vieillir_prets(conn)

    res = services.transferer_pochette(conn, "001", "002", clore_oubli=True)

    assert res["oubli_clos"] is True
    motifs = dict(conn.execute(
        "SELECT id_pret, motif FROM prets ORDER BY id_pret").fetchall())
    # 1 : la boîte rendue, prêt ordinaire ; 2 : le prêt oublié ; 3 : le nouveau.
    assert motifs == {1: "pret", 2: services.MOTIF_OUBLI, 3: "pret"}


def test_les_deux_chemins_ecrivent_la_meme_ligne(conn, vieillir_prets):
    """
    Même clôture des deux côtés : motif, numéro effacé (D5), pochette
    libérée. C'est la garantie que les deux gestes restent comparables.
    """
    services.preter(conn, "001")
    services.preter(conn, "002")
    services.preter(conn, "003")
    vieillir_prets(conn)
    services.repreter(conn, "002")
    services.transferer_pochette(conn, "001", "003", clore_oubli=True)

    oublis = conn.execute(
        "SELECT numero_pochette, date_retour FROM prets WHERE motif = ?",
        (services.MOTIF_OUBLI,),
    ).fetchall()
    assert len(oublis) == 2
    assert all(o["numero_pochette"] is None and o["date_retour"] for o in oublis)


def test_une_sortie_tournoi_oubliee_garde_son_motif(conn):
    """
    Même règle que l'erreur de prêt : `tournoi` est déjà hors statistiques et
    porte une information qu'« oubli » effacerait.
    """
    services.preter(conn, "001")
    services.sortir_tournoi(conn, "002")

    services.transferer_pochette(conn, "001", "002", clore_oubli=True)

    assert _motifs(conn)[1] == "tournoi"


def test_le_re_pret_d_une_boite_disponible_ne_pose_rien(conn):
    """Rien à clore : aucun `oubli` ne doit apparaître."""
    services.repreter(conn, "001")

    assert _motifs(conn) == ["pret"]


def test_un_transfert_ordinaire_ne_pose_pas_le_motif(conn, vieillir_prets):
    services.preter(conn, "001")
    vieillir_prets(conn)

    services.transferer_pochette(conn, "001", "002")

    assert services.MOTIF_OUBLI not in _motifs(conn)


# ---------------------------------------------------------------------------
# 2. Statistiques, liste détaillée, exports
# ---------------------------------------------------------------------------
@pytest.fixture
def un_oubli_et_un_pret(conn, vieillir_prets):
    """
    Un prêt ordinaire d'une heure (001, rendu), et un prêt oublié (002) dont
    l'oubli n'est découvert que dix heures après sa sortie, au re-prêt.
    """
    services.preter(conn, "001")
    services.preter(conn, "002")
    # 30 s de plus qu'une heure : la moyenne passe par `julianday`, dont
    # l'arrondi ferait lire « 59 min » pour une heure pile.
    vieillir_prets(conn, 3630)
    services.rendre(conn, "001")
    vieillir_prets(conn, 9 * 3600)          # 002 sorti depuis dix heures
    services.repreter(conn, "002")
    return conn


def test_un_oubli_compte_comme_un_pret(un_oubli_et_un_pret):
    conn = un_oubli_et_un_pret
    g = services.stats_globales(conn)

    # 001, l'ancien prêt de 002 (oubli) et le nouveau prêt de 002.
    assert g["total_prets"] == 3
    assert g["en_cours"] == 1
    assert g["titres_pretes"] == 2
    assert g["oublis"] == 1
    assert g["erreurs"] == 0
    assert sum(h["n"] for h in services.prets_par_heure(conn)) == 3
    par_titre = {l["reference_titre"]: l["nb_prets"]
                 for l in services.palmares(conn, sens="desc")}
    assert par_titre == {"CATAN": 1, "DIXIT": 2}


def test_un_oubli_est_absent_de_la_duree_moyenne(un_oubli_et_un_pret):
    """
    Le prêt ordinaire a duré une heure ; l'oubli « dix heures ». Compté, il
    porterait la moyenne à 5 h 30. La moyenne doit rester d'une heure.
    """
    g = services.stats_globales(un_oubli_et_un_pret)

    assert g["duree_moyenne"] == services.format_duree(3630)


def test_la_liste_detaillee_dit_pourquoi_la_duree_manque(un_oubli_et_un_pret):
    lignes = services.lister_prets_periode(un_oubli_et_un_pret)

    oubli = [l for l in lignes if l["motif"] == services.MOTIF_OUBLI]
    assert len(oubli) == 1
    # Ni une durée inventée, ni l'heure de la découverte présentée comme un
    # retour.
    assert oubli[0]["retour_local"] == services.RETOUR_NON_SCANNE
    assert oubli[0]["duree_txt"] == services.DUREE_INCONNUE
    assert len(lignes) == 3


def test_les_exports_portent_les_oublis(un_oubli_et_un_pret):
    import openpyxl

    data = services.collecter_stats(un_oubli_et_un_pret)

    classeur = openpyxl.load_workbook(BytesIO(exports.construire_xlsx(data, "tout")))
    synthese = {c[0]: c[1] for c in classeur["Synthèse"].iter_rows(values_only=True)
                if c[0]}
    assert synthese[exports.LIBELLE_OUBLIS] == 1
    textes = [str(v) for f in classeur.worksheets
              for ligne in f.iter_rows(values_only=True) for v in ligne if v]
    assert services.RETOUR_NON_SCANNE in textes
    assert services.DUREE_INCONNUE in textes
    assert exports.construire_pdf(data, "tout")


# ---------------------------------------------------------------------------
# 3. Base d'avant
# ---------------------------------------------------------------------------
def test_une_base_d_avant_garde_ses_lignes_telles_quelles(tmp_path):
    """
    Avant ce lot, un re-prêt laissait `pret` sur l'ancienne ligne. Ces lignes
    ne sont PAS requalifiées : rien ne les distingue d'un vrai prêt long, et
    deviner serait pire que laisser. Rouvrir la base ne doit rien changer,
    ni en base ni dans les chiffres.
    """
    from app import db

    chemin = tmp_path / "avant.db"
    c = sqlite3.connect(chemin)
    c.row_factory = sqlite3.Row
    db.init_db(c)
    c.execute("INSERT INTO titres (reference_titre, nom) VALUES ('CATAN', 'Catan')")
    c.execute("INSERT INTO exemplaires (id_exemplaire, reference_titre) "
              "VALUES ('001', 'CATAN')")
    # Ce qu'écrivait l'ancien `repreter` : l'ancienne ligne close en `pret`.
    c.execute("INSERT INTO prets (id_exemplaire, numero_pochette, date_sortie, "
              "date_retour, motif) VALUES ('001', NULL, "
              "'2026-07-20T08:00:00+00:00', '2026-07-20T18:00:30+00:00', 'pret')")
    c.commit()
    avant = [dict(r) for r in c.execute("SELECT * FROM prets")]
    c.close()

    c = sqlite3.connect(chemin)
    c.row_factory = sqlite3.Row
    db.init_db(c)                           # migrations idempotentes rejouées
    try:
        assert [dict(r) for r in c.execute("SELECT * FROM prets")] == avant
        g = services.stats_globales(c)
        # Les chiffres d'avant : un prêt de dix heures, compté dans la durée.
        assert g["total_prets"] == 1
        assert g["duree_moyenne"] == services.format_duree(10 * 3600 + 30)
        assert g["oublis"] == 0
        assert services.lister_prets_periode(c)[0]["duree_txt"] == \
            services.format_duree(10 * 3600 + 30)
    finally:
        c.close()
