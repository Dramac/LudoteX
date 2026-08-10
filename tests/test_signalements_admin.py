"""
Face ADMINISTRATION du carnet de maintenance (docs/conception-signalements.md
§2 et §7, lot 3 de docs/prompt-impl-signalements.md).

Les services sont déjà couverts par tests/test_signalements.py : ici on
vérifie le CÂBLAGE — garde admin sur toutes les routes, CRUD des catégories
transposé de /admin/rangement, liste et ses filtres, « Marquer traité », les
deux emplacements côte à côte, exports, compteur du tableau de bord.

Fixture sur le patron de tests/test_rangement.py (section administration).
"""

import pytest


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_PATH", str(tmp_path / "test.db"))
    monkeypatch.setenv("TOURNOI_DATABASE_PATH", str(tmp_path / "tournoi.db"))
    monkeypatch.setenv("PLANNING_DATABASE_PATH", str(tmp_path / "planning.db"))
    monkeypatch.setenv("ADMIN_PASSWORD", "secret-admin-123")
    from app import db as _db
    from app.planning import db as pdb
    from app.tournoi import db as tdb

    monkeypatch.setattr(_db, "get_database_path", lambda: tmp_path / "test.db")
    monkeypatch.setattr(tdb, "get_database_path", lambda: tmp_path / "tournoi.db")
    monkeypatch.setattr(pdb, "get_database_path", lambda: tmp_path / "planning.db")
    tdb.init_db()
    pdb.init_db()
    conn = _db.get_connection()
    _db.init_db(conn)
    conn.execute("INSERT INTO titres (reference_titre, nom) VALUES ('CATAN', 'Catan')")
    conn.execute("INSERT INTO titres (reference_titre, nom) VALUES ('DIXIT', 'Dixit')")
    conn.execute("INSERT INTO exemplaires (id_exemplaire, reference_titre) VALUES ('001', 'CATAN')")
    conn.execute("INSERT INTO exemplaires (id_exemplaire, reference_titre) VALUES ('002', 'DIXIT')")
    conn.commit()
    conn.close()

    from fastapi.testclient import TestClient

    from app.main import app

    return TestClient(app)


def _connecter(client):
    client.post("/admin/login", data={"mot_de_passe": "secret-admin-123"})


def _conn():
    from app import db

    return db.get_connection()


def _id_categorie(nom_partiel: str) -> int:
    """Identifiant de la première catégorie dont le nom contient `nom_partiel`."""
    conn = _conn()
    try:
        for c in conn.execute(
            "SELECT id_categorie, nom FROM categories_signalement ORDER BY ordre"
        ):
            if nom_partiel.lower() in c["nom"].lower():
                return c["id_categorie"]
    finally:
        conn.close()
    raise AssertionError(f"Aucune catégorie amorcée ne contient « {nom_partiel} »")


def _creer_signalement(id_exemplaire="001", categorie="pièce", texte=None) -> int:
    from app import services

    conn = _conn()
    try:
        return services.creer_signalement(
            conn, id_exemplaire, _id_categorie(categorie), texte
        )
    finally:
        conn.close()


def _affecter_emplacements(id_exemplaire="001", salle=None, local=None):
    """Renseigne l'un et/ou l'autre des deux emplacements d'une boîte."""
    from app import services

    conn = _conn()
    try:
        if salle is not None:
            services.affecter_emplacement(conn, id_exemplaire, "evenement", salle)
        if local is not None:
            id_local = services.obtenir_ou_creer_emplacement_rangement(conn, local)[0]
            services.affecter_emplacement(conn, id_exemplaire, "local", id_local)
    finally:
        conn.close()


# ===========================================================================
# Garde admin — toutes les routes de ce lot
# ===========================================================================
@pytest.mark.parametrize("chemin", [
    "/admin/categories-signalement",
    "/admin/signalements",
    "/admin/signalements/export.xlsx",
    "/admin/signalements/export.pdf",
])
def test_les_pages_redirigent_sans_session_admin(client, chemin):
    r = client.get(chemin, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/admin"


@pytest.mark.parametrize("chemin", [
    "/admin/categories-signalement",
    "/admin/categories-signalement/1/renommer",
    "/admin/categories-signalement/1/archiver",
    "/admin/categories-signalement/1/reactiver",
    "/admin/categories-signalement/1/supprimer",
    "/admin/categories-signalement/1/monter",
    "/admin/categories-signalement/1/descendre",
    "/admin/signalements/1/traiter",
])
def test_les_actions_redirigent_sans_session_admin(client, chemin):
    r = client.post(chemin, data={"nom": "X"}, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/admin"
    # La garde doit REFUSER, pas seulement rediriger : rien ne doit avoir été
    # écrit au passage.
    conn = _conn()
    try:
        assert conn.execute(
            "SELECT COUNT(*) FROM categories_signalement WHERE nom = 'X'"
        ).fetchone()[0] == 0
    finally:
        conn.close()


# ===========================================================================
# CRUD des catégories (patron /admin/rangement)
# ===========================================================================
def test_page_liste_les_5_categories_amorcees(client):
    _connecter(client)
    r = client.get("/admin/categories-signalement")
    assert r.status_code == 200
    for nom in ["Pièce manquante", "Règle manquante", "Boîte abîmée",
                "Matériel abîmé", "Autre"]:
        assert nom in r.text


def test_page_rappelle_que_la_liste_est_lue_au_comptoir(client):
    """
    La consigne du §2 vit sur l'écran, à l'endroit exact où l'on ajoute une
    entrée — pas dans le wiki.
    """
    _connecter(client)
    r = client.get("/admin/categories-signalement")
    assert "comptoir" in r.text


def test_creer_une_categorie(client):
    _connecter(client)
    r = client.post("/admin/categories-signalement", data={"nom": "Jeton manquant"},
                    follow_redirects=True)
    assert "Catégorie ajoutée." in r.text
    assert "Jeton manquant" in r.text


def test_creer_sans_nom_ne_cree_rien(client):
    _connecter(client)
    conn = _conn()
    try:
        avant = conn.execute("SELECT COUNT(*) FROM categories_signalement").fetchone()[0]
    finally:
        conn.close()
    r = client.post("/admin/categories-signalement", data={"nom": "   "},
                    follow_redirects=True)
    assert "Nom manquant" in r.text
    conn = _conn()
    try:
        assert conn.execute(
            "SELECT COUNT(*) FROM categories_signalement"
        ).fetchone()[0] == avant
    finally:
        conn.close()


def test_renommer_se_repercute_sur_les_signalements_existants(client):
    """
    Référence et non libellé recopié (§4) : renommer corrige aussi
    l'historique, y compris les signalements déjà traités.
    """
    _connecter(client)
    id_cat = _id_categorie("pièce")
    _creer_signalement(categorie="pièce")

    r = client.post(f"/admin/categories-signalement/{id_cat}/renommer",
                    data={"nom": "Pièce égarée"}, follow_redirects=True)
    assert "Catégorie renommée." in r.text

    from app import services
    conn = _conn()
    try:
        noms = {s["categorie_nom"] for s in services.lister_signalements(conn, "tous")}
    finally:
        conn.close()
    assert noms == {"Pièce égarée"}


def test_renommer_sans_nom_ne_modifie_rien(client):
    _connecter(client)
    id_cat = _id_categorie("pièce")
    r = client.post(f"/admin/categories-signalement/{id_cat}/renommer",
                    data={"nom": ""}, follow_redirects=True)
    assert "Nom manquant" in r.text
    assert "Pièce manquante" in r.text


def test_archiver_puis_reactiver(client):
    _connecter(client)
    id_cat = _id_categorie("Autre")

    r = client.post(f"/admin/categories-signalement/{id_cat}/archiver",
                    follow_redirects=True)
    assert "Catégorie archivée" in r.text
    assert "Archivée" in r.text

    r2 = client.post(f"/admin/categories-signalement/{id_cat}/reactiver",
                     follow_redirects=True)
    assert "Catégorie réactivée." in r2.text


def test_archiver_ne_supprime_aucun_signalement(client):
    """L'archivage est doux : la FK est sans cascade (§4)."""
    _connecter(client)
    id_cat = _id_categorie("boîte")
    _creer_signalement(categorie="boîte")

    client.post(f"/admin/categories-signalement/{id_cat}/archiver")

    from app import services
    conn = _conn()
    try:
        signalements = services.lister_signalements(conn, "tous")
    finally:
        conn.close()
    assert [s["categorie_nom"] for s in signalements] == ["Boîte abîmée"]


def test_supprimer_une_categorie_sans_usage(client):
    _connecter(client)
    id_cat = _id_categorie("Autre")
    r = client.post(f"/admin/categories-signalement/{id_cat}/supprimer",
                    follow_redirects=True)
    assert "Catégorie supprimée définitivement." in r.text
    conn = _conn()
    try:
        assert conn.execute(
            "SELECT COUNT(*) FROM categories_signalement WHERE id_categorie = ?",
            (id_cat,),
        ).fetchone()[0] == 0
    finally:
        conn.close()


def test_supprimer_est_refuse_sous_usage(client):
    _connecter(client)
    id_cat = _id_categorie("règle")
    _creer_signalement(categorie="règle")

    r = client.post(f"/admin/categories-signalement/{id_cat}/supprimer",
                    follow_redirects=True)
    assert "Suppression refusée" in r.text
    assert "Règle manquante" in r.text  # toujours dans la liste, rien n'a bougé


def test_monter_et_descendre_changent_l_ordre(client):
    _connecter(client)
    id_second = _id_categorie("règle")  # 2e du seed

    client.post(f"/admin/categories-signalement/{id_second}/monter")
    conn = _conn()
    try:
        premiers = [
            r["nom"] for r in conn.execute(
                "SELECT nom FROM categories_signalement ORDER BY ordre, nom"
            )
        ]
    finally:
        conn.close()
    assert premiers[0] == "Règle manquante"

    client.post(f"/admin/categories-signalement/{id_second}/descendre")
    conn = _conn()
    try:
        premiers = [
            r["nom"] for r in conn.execute(
                "SELECT nom FROM categories_signalement ORDER BY ordre, nom"
            )
        ]
    finally:
        conn.close()
    assert premiers[0] == "Pièce manquante"


# ===========================================================================
# La liste des signalements (§7)
# ===========================================================================
def test_liste_vide_le_dit_sans_tableau(client):
    _connecter(client)
    r = client.get("/admin/signalements")
    assert r.status_code == 200
    assert "Aucun signalement en attente" in r.text


def test_liste_affiche_jeu_boite_categorie_detail_et_date(client):
    _connecter(client)
    _creer_signalement("001", "pièce", "il manque un dé rouge")

    r = client.get("/admin/signalements")
    assert "Catan" in r.text
    assert "001" in r.text
    assert "Pièce manquante" in r.text
    assert "il manque un dé rouge" in r.text


def test_compteur_d_ouverts_en_tete(client):
    _connecter(client)
    _creer_signalement("001", "pièce")
    _creer_signalement("002", "boîte")

    r = client.get("/admin/signalements")
    assert "2 signalements ouverts" in r.text


def test_compteur_au_singulier(client):
    """Le compteur s'accorde (global Jinja `pluriel`, fiche Q2)."""
    _connecter(client)
    _creer_signalement("001", "pièce")

    r = client.get("/admin/signalements")
    assert "1 signalement ouvert" in r.text


def test_les_deux_emplacements_cote_a_cote(client):
    """
    §7 : les DEUX emplacements, pas celui du contexte actif — la liste se
    traite après l'événement, quand le contexte est repassé en « local ».
    """
    _connecter(client)
    _affecter_emplacements("001", salle="Table 4", local="Étagère 3")
    _creer_signalement("001", "pièce")

    r = client.get("/admin/signalements")
    assert "Table 4" in r.text
    assert "Étagère 3" in r.text


def test_un_seul_emplacement_renseigne_n_affiche_que_lui(client):
    _connecter(client)
    _affecter_emplacements("001", salle="Table 4")
    _creer_signalement("001", "pièce")

    r = client.get("/admin/signalements")
    assert "Table 4" in r.text
    assert "Local :" not in r.text


def test_aucun_emplacement_n_affiche_rien(client):
    """Jamais de « non renseigné » : une case vide, et c'est tout."""
    _connecter(client)
    _creer_signalement("001", "pièce")

    r = client.get("/admin/signalements")
    assert "Salle :" not in r.text
    assert "Local :" not in r.text
    assert "non renseigné" not in r.text.lower()


def test_filtre_par_etat(client):
    # Détails volontairement distinctifs : `base.html` embarque un script qui
    # contient des mots courants (« premier », « second »…), et une assertion
    # d'ABSENCE dessus passerait au vert pour de mauvaises raisons.
    _connecter(client)
    id_a = _creer_signalement("001", "pièce", "detail-alpha")
    _creer_signalement("002", "boîte", "detail-beta")

    from app import services
    conn = _conn()
    try:
        services.traiter_signalement(conn, id_a)
    finally:
        conn.close()

    ouverts = client.get("/admin/signalements")
    assert "detail-beta" in ouverts.text and "detail-alpha" not in ouverts.text

    traites = client.get("/admin/signalements?etat=traites")
    assert "detail-alpha" in traites.text and "detail-beta" not in traites.text

    tous = client.get("/admin/signalements?etat=tous")
    assert "detail-alpha" in tous.text and "detail-beta" in tous.text


def test_filtre_par_categorie(client):
    _connecter(client)
    _creer_signalement("001", "pièce", "detail-alpha")
    _creer_signalement("002", "boîte", "detail-beta")

    r = client.get(f"/admin/signalements?categorie={_id_categorie('boîte')}")
    assert "detail-beta" in r.text
    assert "detail-alpha" not in r.text


def test_filtre_de_categorie_forge_est_ignore(client):
    """Jamais bloquant : un identifiant inconnu retombe sur « toutes »."""
    _connecter(client)
    _creer_signalement("001", "pièce", "un dé")

    r = client.get("/admin/signalements?categorie=99999")
    assert r.status_code == 200
    assert "un dé" in r.text


def test_etat_forge_retombe_sur_ouverts(client):
    _connecter(client)
    _creer_signalement("001", "pièce", "un dé")

    r = client.get("/admin/signalements?etat=nimporte-quoi")
    assert r.status_code == 200
    assert "un dé" in r.text


def test_puces_de_retrait_des_filtres(client):
    _connecter(client)
    id_cat = _id_categorie("boîte")
    _creer_signalement("002", "boîte")

    r = client.get(f"/admin/signalements?etat=tous&categorie={id_cat}")
    assert "Filtres actifs" in r.text
    # Retirer l'état seul conserve la catégorie, et réciproquement.
    assert f"/admin/signalements?categorie={id_cat}" in r.text
    assert "/admin/signalements?etat=tous" in r.text


def test_marquer_traite(client):
    _connecter(client)
    id_s = _creer_signalement("001", "pièce", "un dé")

    r = client.post(f"/admin/signalements/{id_s}/traiter",
                    data={"etat": "ouverts", "categorie": ""},
                    follow_redirects=True)
    assert "Signalement marqué traité." in r.text
    assert "Aucun signalement en attente" in r.text
    assert "0 signalement ouvert" in r.text


def test_marquer_traite_conserve_les_filtres(client):
    _connecter(client)
    id_cat = _id_categorie("boîte")
    id_s = _creer_signalement("002", "boîte")

    r = client.post(f"/admin/signalements/{id_s}/traiter",
                    data={"etat": "tous", "categorie": str(id_cat)},
                    follow_redirects=False)
    assert r.status_code == 303
    destination = r.headers["location"]
    assert "etat=tous" in destination and f"categorie={id_cat}" in destination


def test_marquer_traite_est_idempotent(client):
    """Deux appuis (double-clic, page rechargée) : même résultat, aucune erreur."""
    _connecter(client)
    id_s = _creer_signalement("001", "pièce")

    client.post(f"/admin/signalements/{id_s}/traiter", data={"etat": "ouverts"})
    conn = _conn()
    try:
        premier = conn.execute(
            "SELECT traite_le FROM signalements WHERE id_signalement = ?", (id_s,)
        ).fetchone()["traite_le"]
    finally:
        conn.close()

    r = client.post(f"/admin/signalements/{id_s}/traiter", data={"etat": "ouverts"},
                    follow_redirects=True)
    assert r.status_code == 200
    conn = _conn()
    try:
        second = conn.execute(
            "SELECT traite_le FROM signalements WHERE id_signalement = ?", (id_s,)
        ).fetchone()["traite_le"]
    finally:
        conn.close()
    assert second == premier  # la date de traitement n'est pas réécrite


def test_une_ligne_traitee_n_a_pas_de_bouton(client):
    _connecter(client)
    id_s = _creer_signalement("001", "pièce")
    client.post(f"/admin/signalements/{id_s}/traiter", data={"etat": "ouverts"})

    # On vise le FORMULAIRE, pas le libellé : « Marquer traité » figure aussi
    # dans le bloc d'aide replié de la page.
    r = client.get("/admin/signalements?etat=traites")
    assert "/traiter" not in r.text
    assert "Traité" in r.text


def test_la_page_est_large(client):
    """
    Piège 5 du lot 3 : un tableau dense a besoin de `contenu-large`, sinon le
    max-width de 540 px le bride sur ordinateur (erreur déjà commise six fois).
    """
    _connecter(client)
    r = client.get("/admin/signalements")
    assert "contenu-large" in r.text


# ===========================================================================
# Exports (§7) — réservés à cette page, jamais /stats
# ===========================================================================
def _feuille(contenu: bytes):
    import io

    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(contenu)).active


def test_export_xlsx_contient_les_colonnes_et_la_ligne(client):
    _connecter(client)
    _affecter_emplacements("001", salle="Table 4", local="Étagère 3")
    _creer_signalement("001", "pièce", "il manque un dé rouge")

    r = client.get("/admin/signalements/export.xlsx")
    assert r.status_code == 200
    assert "signalements.xlsx" in r.headers["content-disposition"]

    ws = _feuille(r.content)
    assert ws.title == "Signalements"  # et surtout pas « Catalogue »
    entetes = [c.value for c in ws[1]]
    assert entetes == [
        "Jeu", "Boîte", "Catégorie", "Détail", "Signalé le", "Traité le",
        "Emplacement salle", "Emplacement local",
    ]
    ligne = [c.value for c in ws[2]]
    assert ligne[0] == "Catan"
    assert ligne[1] == "001"
    assert ligne[2] == "Pièce manquante"
    assert ligne[3] == "il manque un dé rouge"
    assert ligne[6] == "Table 4"
    assert ligne[7] == "Étagère 3"


def test_export_xlsx_respecte_les_filtres(client):
    _connecter(client)
    _creer_signalement("001", "pièce", "detail-alpha")
    _creer_signalement("002", "boîte", "detail-beta")

    r = client.get(f"/admin/signalements/export.xlsx?categorie={_id_categorie('boîte')}")
    ws = _feuille(r.content)
    details = [ws.cell(row=i, column=4).value for i in range(2, ws.max_row + 1)]
    assert details == ["detail-beta"]


def test_export_xlsx_sans_signalement_garde_ses_entetes(client):
    _connecter(client)
    r = client.get("/admin/signalements/export.xlsx")
    ws = _feuille(r.content)
    assert ws.max_row == 1
    assert ws.cell(row=1, column=1).value == "Jeu"


def test_export_pdf(client, monkeypatch):
    """
    Le contenu d'un PDF reportlab est compressé, donc illisible dans les
    octets bruts : on intercepte `Table` pour capturer ce qui lui est
    réellement transmis (patron du volet 1 de la fiche D5).
    """
    _connecter(client)
    _affecter_emplacements("001", salle="Table 4")
    _creer_signalement("001", "pièce", "il manque un dé rouge")

    import reportlab.platypus as platypus

    captures = []
    vraie_table = platypus.Table

    def _table_espionne(donnees, *args, **kwargs):
        captures.append(donnees)
        return vraie_table(donnees, *args, **kwargs)

    monkeypatch.setattr(platypus, "Table", _table_espionne)

    r = client.get("/admin/signalements/export.pdf")
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.content[:4] == b"%PDF"
    assert "signalements.pdf" in r.headers["content-disposition"]

    assert captures, "aucun tableau transmis à la mise en page"
    texte = " ".join(
        cellule.getPlainText() for ligne in captures[0] for cellule in ligne
    )
    assert "Catan" in texte
    assert "il manque un dé rouge" in texte
    assert "Table 4" in texte


def test_export_pdf_sans_signalement_ne_plante_pas(client):
    """Jamais bloquant : une liste vide produit un PDF, pas une erreur."""
    _connecter(client)
    r = client.get("/admin/signalements/export.pdf")
    assert r.status_code == 200 and r.content[:4] == b"%PDF"


def test_rien_de_tout_cela_ne_sort_par_stats(client):
    """
    Piège 2 du lot 3 : /stats est PUBLIQUE. Le carnet de maintenance nomme des
    boîtes abîmées et porte du texte libre — il ne doit apparaître ni sur la
    page, ni dans ses deux exports (précédent D5 sur le numéro de pochette).
    """
    _connecter(client)
    _creer_signalement("001", "pièce", "texte-libre-distinctif")

    page = client.get("/stats")
    assert page.status_code == 200
    assert "texte-libre-distinctif" not in page.text
    assert "signalement" not in page.text.lower()

    xlsx = client.get("/stats/export.xlsx")
    assert b"texte-libre-distinctif" not in xlsx.content

    pdf = client.get("/stats/export.pdf")
    assert b"texte-libre-distinctif" not in pdf.content
