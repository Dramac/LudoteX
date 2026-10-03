"""
Neutralisation des formules dans les exports tableur (SEC-06, lot 14).

Une cellule qui commence par « = », « + », « - », « @ », une tabulation ou un
retour chariot est interprétée par un tableur. Le détail d'un signalement est
tapé au comptoir, le nom d'un bénévole dans le questionnaire public : ces
valeurs ne viennent plus seulement de l'admin.

Ce qui est vérifié :
- les six débuts dangereux, dans les deux fonctions d'export génériques
  (`catalogue_csv`, `tableau_xlsx`) ;
- en Excel, la valeur reste EXACTE (aucune apostrophe dans la cellule) et
  typée texte ; en CSV, l'apostrophe est posée — c'est la seule défense
  possible dans un format sans type ;
- l'aller-retour réel du bureau : export CSV du catalogue → import, sans
  qu'une apostrophe ne s'ajoute au nom des jeux ;
- les deux autres classeurs qui écrivent des données saisies (statistiques,
  planning) passent par la même fonction.
"""

from __future__ import annotations

import csv
import io
from io import BytesIO

import openpyxl
import pytest

from app import exports

MOT_DE_PASSE = "secret-admin-formules"

# Les six débuts dangereux, chacun suivi d'une formule plausible.
DANGEREUSES = [f"{debut}HYPERLINK(\"http://exemple.invalid\")" for debut in exports.DEBUTS_FORMULE]


def _csv_cases(contenu: bytes) -> list[dict]:
    texte = contenu.decode("utf-8-sig")
    return list(csv.DictReader(io.StringIO(texte, newline=""), delimiter=";"))


@pytest.mark.parametrize("valeur", DANGEREUSES, ids=[repr(d) for d in exports.DEBUTS_FORMULE])
def test_csv_prefixe_les_debuts_de_formule(valeur):
    contenu = exports.catalogue_csv(["Nom jeu"], [{"Nom jeu": valeur}])
    assert _csv_cases(contenu)[0]["Nom jeu"] == "'" + valeur


@pytest.mark.parametrize("valeur", DANGEREUSES, ids=[repr(d) for d in exports.DEBUTS_FORMULE])
def test_xlsx_type_texte_sans_apostrophe(valeur):
    contenu = exports.tableau_xlsx(["Détail"], [{"Détail": valeur}], "Signalements")
    cellule = openpyxl.load_workbook(BytesIO(contenu)).active["A2"]
    assert cellule.data_type == "s"          # jamais « f » (formule)
    # Valeur exacte : pas d'apostrophe visible. Seule nuance, étrangère à la
    # neutralisation : le XML normalise un retour chariot en saut de ligne.
    assert cellule.value == valeur.replace("\r", "\n")
    assert cellule.quotePrefix is True       # Excel ne la réinterprète pas si on la retape


def test_valeurs_ordinaires_inchangees():
    lignes = [{"A": "Catan", "B": "2 - 4", "C": "l'Âge de pierre", "D": ""}]
    assert _csv_cases(exports.catalogue_csv(list("ABCD"), lignes))[0] == lignes[0]
    feuille = openpyxl.load_workbook(BytesIO(exports.tableau_xlsx(list("ABCD"), lignes))).active
    assert [c.value for c in feuille[2]] == ["Catan", "2 - 4", "l'Âge de pierre", None]
    assert not any(c.quotePrefix for c in feuille[2])


def test_retirer_neutralisation_est_l_inverse_exact():
    for valeur in DANGEREUSES + ["Catan", "'Twas the night", "'", "", "-"]:
        assert exports.retirer_neutralisation_csv(exports.neutraliser_csv(valeur)) == valeur
    # Une apostrophe légitime, non suivie d'un début de formule, est préservée.
    assert exports.retirer_neutralisation_csv("'Twas") == "'Twas"
    # Une ligne trop longue range son surplus en liste : rendue telle quelle.
    assert exports.retirer_neutralisation_csv(["x"]) == ["x"]


# ---------------------------------------------------------------------------
# Aller-retour export → import du catalogue (le va-et-vient réel du bureau)
# ---------------------------------------------------------------------------
@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_PATH", str(tmp_path / "test.db"))
    monkeypatch.setenv("TOURNOI_DATABASE_PATH", str(tmp_path / "tournoi.db"))
    monkeypatch.setenv("PLANNING_DATABASE_PATH", str(tmp_path / "planning.db"))
    from app import admin_auth, db
    from app.planning import db as pdb
    from app.tournoi import db as tdb

    monkeypatch.setattr(db, "get_database_path", lambda: tmp_path / "test.db")
    monkeypatch.setattr(tdb, "get_database_path", lambda: tmp_path / "tournoi.db")
    monkeypatch.setattr(pdb, "get_database_path", lambda: tmp_path / "planning.db")
    monkeypatch.setattr(admin_auth, "_sessions", {})
    monkeypatch.setenv("ADMIN_PASSWORD", MOT_DE_PASSE)
    tdb.init_db()
    pdb.init_db()
    conn = db.get_connection()
    db.init_db(conn)
    # Références calculées comme l'import les calcule (`slug_titre`) : le
    # réimport retombe sur les MÊMES lignes au lieu d'en créer d'autres.
    from app.services import slug_titre

    jeux = [("001", "=SOMME(A1:A2)", "- pour deux joueurs", "@editeur"),
            ("002", "+1 jeu", "Rien à signaler", "Ordinaire")]
    for id_ex, nom, descriptif, editeur in jeux:
        conn.execute(
            "INSERT INTO titres (reference_titre, nom, descriptif, editeur) VALUES (?, ?, ?, ?)",
            (slug_titre(nom), nom, descriptif, editeur))
        conn.execute("INSERT INTO exemplaires (id_exemplaire, reference_titre) VALUES (?, ?)",
                     (id_ex, slug_titre(nom)))
    conn.commit()
    conn.close()

    from fastapi.testclient import TestClient

    from app.main import app

    c = TestClient(app)
    c.post("/admin/login", data={"mot_de_passe": MOT_DE_PASSE}, follow_redirects=False)
    return c


def _titres():
    from app.db import get_connection

    conn = get_connection()
    try:
        return {r["reference_titre"]: (r["nom"], r["descriptif"], r["editeur"])
                for r in conn.execute("SELECT reference_titre, nom, descriptif, editeur FROM titres")}
    finally:
        conn.close()


def test_aller_retour_export_import_du_catalogue(client, tmp_path):
    from scripts.import_csv import importer

    avant = _titres()
    r = client.get("/admin/donnees/export.csv")
    assert r.status_code == 200
    # L'export est bien neutralisé…
    noms = [l["Nom jeu"] for l in _csv_cases(r.content)]
    assert "'=SOMME(A1:A2)" in noms and "'+1 jeu" in noms
    # … et le réimporter tel quel ne change rien : aucune apostrophe ne s'ajoute.
    fichier = tmp_path / "catalogue.csv"
    fichier.write_bytes(r.content)
    importer(fichier)
    assert _titres() == avant


def test_export_xlsx_du_catalogue_garde_les_valeurs(client):
    r = client.get("/admin/donnees/export.xlsx")
    feuille = openpyxl.load_workbook(BytesIO(r.content)).active
    valeurs = {c.value for ligne in feuille.iter_rows(min_row=2) for c in ligne}
    assert {"=SOMME(A1:A2)", "- pour deux joueurs", "@editeur", "+1 jeu"} <= valeurs
    assert all(c.data_type != "f" for ligne in feuille.iter_rows() for c in ligne)


# ---------------------------------------------------------------------------
# Les deux autres classeurs qui écrivent des données saisies
# ---------------------------------------------------------------------------
def test_statistiques_xlsx_nom_de_jeu_jamais_formule():
    data = {
        "globales": {"total_prets": 1, "en_cours": 0, "titres_pretes": 1, "nb_titres": 1},
        "metrique": "total",
        "plus": [{"nom": "=1+1", "nb_prets": 1, "nb_exemplaires": 1, "par_exemplaire": 1}],
        "moins": [],
        "prets": [{"nom": "@jeu", "id_exemplaire": "-001", "sortie_locale": "01/01/2026 10:00",
                   "retour_local": None, "duree_txt": "—"}],
    }
    classeur = openpyxl.load_workbook(BytesIO(exports.construire_xlsx(data, "tout")))
    assert classeur["Palmarès"]["A5"].value == "=1+1"
    assert classeur["Palmarès"]["A5"].data_type == "s"
    assert [classeur["Détail"]["A2"].value, classeur["Détail"]["B2"].value] == ["@jeu", "-001"]
    assert all(c.data_type != "f" for f in classeur for l in f.iter_rows() for c in l)


def test_planning_xlsx_nom_de_benevole_jamais_formule():
    from app.planning import exports as pexports

    creneau = {"debut": "2026-01-01T09:00:00+00:00", "fin": "2026-01-01T11:00:00+00:00",
               "libelle": "=Montage", "libelle_jour": "Samedi"}
    grille = {
        "postes": [{"nom": "+Accueil"}],
        "jours": [{"libelle": "Samedi", "creneaux": [
            {"creneau": creneau,
             "cases": [{"nb_requis": 1, "affectations": [{"nom": "=HYPERLINK(1)"}]}]},
        ]}],
        "taches": [{"creneau": creneau, "affectations": [{"nom": "@Camille"}]}],
    }
    classeur = openpyxl.load_workbook(BytesIO(pexports.construire_xlsx(grille, "Édition")))
    samedi, taches = classeur["Samedi"], classeur["Tâches"]
    assert [samedi["B2"].value, samedi["B3"].value] == ["+Accueil", "=HYPERLINK(1)"]
    assert [taches["A3"].value, taches["C3"].value] == ["=Montage", "@Camille"]
    assert all(c.data_type != "f" for f in classeur for l in f.iter_rows() for c in l)
